#!/usr/bin/env python3
"""Structural checks on the built handbook."""
import re
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).parent
t = (ROOT / "combined.md").read_text(encoding="utf-8")
H = [m.group(2) for m in re.finditer(r"^(#{1,3}) (.*)$", t, re.M)]


def nums(prefix):
    return sorted({int(m.group(1)) for h in H for m in [re.match(prefix + r" (\d+)", h)] if m})


print("phases 0-17 all present:", nums("Phase") == list(range(18)))
print("weeks missing:", [w for w in range(1, 25) if w not in nums("Week")])
print("projects found:", nums("Project"))
print("oembed URLs left in appendix:", len(re.findall(r"oembed\?url=https", t.split("# Appendix: Source List")[1])))
d = Document(str(ROOT / "Full-Stack-to-AI-Engineer-Roadmap.docx"))
x = zipfile.ZipFile(ROOT / "Full-Stack-to-AI-Engineer-Roadmap.docx").read("word/document.xml").decode()
print("docx reopens: paragraphs", len(d.paragraphs), "tables", len(d.tables),
      "bookmarks", x.count("w:bookmarkStart"), "internal links", x.count("w:anchor"))
print("checkbox items:", t.count("☐"))
pdf = (ROOT / "Full-Stack-to-AI-Engineer-Roadmap.pdf").read_bytes()
print("pdf pages:", max(map(int, re.findall(rb"/Count (\d+)", pdf))))
# resource tables must have 7 columns
bad = 0
for blk in re.findall(r"\| Topic \| Subtopics[^\n]*", t):
    if len(blk.strip().strip("|").split("|")) != 7:
        bad += 1
print("resource tables with wrong column count:", bad, "of", len(re.findall(r"\| Topic \| Subtopics", t)))
print("emoji/unsupported chars:", len(re.findall(r"[\U0001F300-\U0001FAFF]", t)))
