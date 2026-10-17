#!/usr/bin/env python3
"""Combine parts/*.md into one handbook, then emit DOCX (python-docx) and PDF (headless Chrome)."""
import html
import re
import subprocess
import sys
from pathlib import Path

import markdown
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
PARTS = ROOT / "parts"
TITLE = "The Complete Full-Stack to AI Engineer Roadmap: Fundamentals, Python, FastAPI, GenAI, Agentic AI and Cloud"
SUBTITLE = "A 24-week study handbook for a working React / Next.js developer in India"
DATE = "Research date: 9 October 2026"
OUT_DOCX = ROOT / "Full-Stack-to-AI-Engineer-Roadmap.docx"
OUT_PDF = ROOT / "Full-Stack-to-AI-Engineer-Roadmap.pdf"
OUT_MD = ROOT / "combined.md"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
ORDER = ["p00", "p01", "p02", "p03", "p04", "p05", "p06", "p07a", "p07b", "p07c"]

FENCE = re.compile(r"^\s*```")
HEAD = re.compile(r"^(#{1,6})\s+(.*?)\s*$")


# ---------------------------------------------------------------- combine
def split_sources(text):
    """Remove '## Sources' sections; return (text_without, [urls_lines])."""
    out, srcs, lines, i, in_code = [], [], text.split("\n"), 0, False
    while i < len(lines):
        ln = lines[i]
        if FENCE.match(ln):
            in_code = not in_code
        m = None if in_code else HEAD.match(ln)
        if m and len(m.group(1)) == 2 and m.group(2).strip().lower().startswith("sources"):
            i += 1
            while i < len(lines):
                h = HEAD.match(lines[i])
                if h and len(h.group(1)) <= 2:
                    break
                if lines[i].strip():
                    srcs.append(lines[i].strip())
                i += 1
            continue
        out.append(ln)
        i += 1
    return "\n".join(out), srcs


def combine():
    chunks, source_groups = [], []
    for key in ORDER:
        for f in sorted(PARTS.glob(f"{key}_*.md")):
            text = f.read_text(encoding="utf-8").strip() + "\n"
            text, srcs = split_sources(text)
            chunks.append(text)
            title = next((HEAD.match(l).group(2) for l in text.split("\n") if HEAD.match(l)), f.stem)
            if srcs:
                source_groups.append((title, srcs))
    seen, appendix = set(), ["# Appendix: Source List", "",
                             "Every URL below was opened or checked by the authors on the research date (9 October 2026). "
                             "Grouped by the part of the handbook that uses it.", ""]
    for title, srcs in source_groups:
        rows = []
        for s in srcs:
            s = re.sub(r"https://www\.youtube\.com/oembed\?url=(https://www\.youtube\.com/watch\?v=[\w-]+)&format=json( \(via WebFetch\))?", r"\1", s)
            urls = re.findall(r"https?://[^\s)>\]]+", s)
            if urls and all(u in seen for u in urls):
                continue
            seen.update(urls)
            rows.append(s if s.startswith(("-", "*")) else "- " + s)
        if rows:
            appendix += ["## " + title, ""] + rows + [""]
    chunks.append("\n".join(appendix))
    combined = "\n\n".join(chunks)
    OUT_MD.write_text(combined, encoding="utf-8")
    return combined


# ---------------------------------------------------------------- shared parsing
def normalise_checks(text):
    text = re.sub(r"^(\s*)([-*])\s+\[ \]\s+", r"\1\2 ☐ ", text, flags=re.M)
    text = re.sub(r"^(\s*)([-*])\s+\[[xX]\]\s+", r"\1\2 ☑ ", text, flags=re.M)
    return text


def headings(text):
    out, in_code = [], False
    for ln in text.split("\n"):
        if FENCE.match(ln):
            in_code = not in_code
        m = None if in_code else HEAD.match(ln)
        if m:
            out.append((len(m.group(1)), m.group(2).strip()))
    return out


# ---------------------------------------------------------------- HTML / PDF
CSS = """
@page { size: A4; margin: 16mm 14mm 18mm 14mm; @bottom-center { content: counter(page); font: 9pt Helvetica, Arial, sans-serif; color: #666; } }
body { font: 10pt/1.5 -apple-system, "Segoe UI", Helvetica, Arial, sans-serif; color: #1c1c1c; }
h1 { font-size: 21pt; color: #12355b; border-bottom: 2px solid #12355b; padding-bottom: 4px; page-break-before: always; margin-top: 0; }
h2 { font-size: 15pt; color: #1b4f8a; margin-top: 20px; }
h3 { font-size: 12pt; color: #234; margin-top: 16px; }
h4 { font-size: 10.5pt; color: #345; margin-top: 12px; }
h1, h2, h3, h4 { page-break-after: avoid; }
a { color: #0b5cc2; text-decoration: none; word-break: break-word; }
code { font: 8.8pt Menlo, Consolas, monospace; background: #f1f3f5; padding: 0 3px; border-radius: 3px; }
pre { background: #f6f8fa; border: 1px solid #d8dee4; border-radius: 4px; padding: 8px 10px; white-space: pre-wrap; word-break: break-word; page-break-inside: avoid; }
pre code { background: none; padding: 0; font-size: 8.4pt; }
table { border-collapse: collapse; width: 100%; margin: 8px 0 12px; font-size: 8pt; table-layout: auto; }
th { background: #12355b; color: #fff; text-align: left; padding: 4px 5px; }
td { border: 1px solid #c9d1d9; padding: 3px 5px; vertical-align: top; word-break: break-word; }
tr:nth-child(even) td { background: #f7f9fb; }
tr { page-break-inside: avoid; }
blockquote { margin: 10px 0; padding: 6px 12px; background: #fff7e0; border-left: 4px solid #e0a800; page-break-inside: avoid; }
blockquote p { margin: 3px 0; }
ul, ol { padding-left: 20px; }
li { margin: 2px 0; }
.cover { height: 240mm; display: flex; flex-direction: column; justify-content: center; }
.cover h1 { border: none; font-size: 30pt; page-break-before: avoid; }
.cover p { font-size: 13pt; color: #444; }
.toc h1 { page-break-before: always; }
.toc ul { list-style: none; padding-left: 0; }
.toc li.l2 { margin-left: 14px; font-size: 9.5pt; }
.toc li.l3 { margin-left: 28px; font-size: 8.8pt; color: #555; }
.toc li.l1 { margin-top: 6px; font-weight: 600; }
"""


def build_html(combined):
    text = normalise_checks(combined)
    # doubled indent so 2-space nested lists render under python-markdown
    fixed, in_code = [], False
    for ln in text.split("\n"):
        if FENCE.match(ln):
            in_code = not in_code
        if not in_code:
            m = re.match(r"^(\s+)([-*]|\d+\.)\s", ln)
            if m:
                ln = m.group(1) * 2 + ln[len(m.group(1)):]
        fixed.append(ln)
    text = "\n".join(fixed)
    body = markdown.markdown(text, extensions=["tables", "fenced_code", "toc", "sane_lists"],
                             extension_configs={"toc": {"permalink": False}})
    # open external links in new context is irrelevant for PDF; build TOC from rendered headings
    toc_items = []
    for m in re.finditer(r'<h([123]) id="([^"]+)">(.*?)</h\1>', body):
        lvl, hid, label = int(m.group(1)), m.group(2), re.sub(r"<[^>]+>", "", m.group(3))
        if lvl == 3 and not re.match(r"Week \d+", label):
            continue
        toc_items.append(f'<li class="l{lvl}"><a href="#{hid}">{label}</a></li>')
    toc = '<div class="toc"><h1 id="contents">Contents</h1><ul>' + "".join(toc_items) + "</ul></div>"
    cover = (f'<div class="cover"><h1>{html.escape(TITLE)}</h1><p>{html.escape(SUBTITLE)}</p>'
             f'<p>{html.escape(DATE)}</p></div>')
    return f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{html.escape(TITLE)}</title><style>{CSS}</style></head><body>{cover}{toc}{body}</body></html>'


def build_pdf(combined):
    page = ROOT / "handbook.html"
    page.write_text(build_html(combined), encoding="utf-8")
    cmd = [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--generate-pdf-document-outline",
           f"--print-to-pdf={OUT_PDF}", page.as_uri()]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    if not OUT_PDF.exists():
        print("PDF failed:", r.stderr[-800:])
        return False
    return True


# ---------------------------------------------------------------- DOCX
TOK = re.compile(r"(\[([^\]]+)\]\((https?://[^)\s]+)\))|(\*\*(.+?)\*\*)|(`([^`]+)`)|(\*([^*\s][^*]*?)\*)")


def inline_tokens(text, bold=False, italic=False, code=False):
    pos = 0
    for m in TOK.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], dict(bold=bold, italic=italic, code=code), None
        if m.group(1):
            for t, f, _ in inline_tokens(m.group(2), bold, italic, code):
                yield t, f, m.group(3)
        elif m.group(4):
            yield from inline_tokens(m.group(5), True, italic, code)
        elif m.group(6):
            yield m.group(7), dict(bold=bold, italic=italic, code=True), None
        else:
            yield from inline_tokens(m.group(9), bold, True, code)
        pos = m.end()
    if pos < len(text):
        yield text[pos:], dict(bold=bold, italic=italic, code=code), None


def style_run(run, f, size=None, color=None):
    run.bold = f["bold"] or None
    run.italic = f["italic"] or None
    if f["code"]:
        run.font.name = "Menlo"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Menlo")
        run.font.size = Pt((size or 10) - 1)
        run.font.color.rgb = RGBColor(0x9A, 0x1B, 0x4B)
    if size and not f["code"]:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_hyperlink(par, url, text, f, size=None):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = par.add_run(text)
    style_run(run, f, size)
    run.font.color.rgb = RGBColor(0x0B, 0x5C, 0xC2)
    run.font.underline = True
    par._p.remove(run._r)
    h.append(run._r)
    par._p.append(h)


def add_inline(par, text, size=None, base_bold=False, color=None):
    for t, f, url in inline_tokens(text):
        if base_bold:
            f["bold"] = True
        if url:
            add_hyperlink(par, url, t, f, size)
        else:
            run = par.add_run(t)
            style_run(run, f, size, color)


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def par_border_left(par, color="E0A800"):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "24"), ("space", "6"), ("color", color)):
        left.set(qn("w:" + k), v)
    b.append(left)
    pPr.append(b)


def bookmark(par, name, idx):
    s, e = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    s.set(qn("w:id"), str(idx))
    s.set(qn("w:name"), name)
    e.set(qn("w:id"), str(idx))
    par._p.insert(0, s)
    par._p.append(e)


def internal_link(par, anchor, text, size=10):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x0B, 0x5C, 0xC2)
    par._p.remove(run._r)
    h.append(run._r)
    par._p.append(h)


def add_field(par, instr):
    for kind, txt in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)


def split_row(line):
    line = line.strip().replace("\\|", "\u0001")
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("\u0001", "|") for c in line.split("|")]


def set_cell_bg(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def build_docx(combined):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.7))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"].paragraph_format.space_after = Pt(5)
    sizes = {1: 22, 2: 16, 3: 13, 4: 11}
    for lvl, sz in sizes.items():
        s = st[f"Heading {lvl}"]
        s.font.name = "Calibri"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0x12, 0x35, 0x5B) if lvl == 1 else RGBColor(0x1B, 0x4F, 0x8A)
        s.paragraph_format.space_before = Pt(14 if lvl > 1 else 0)
        s.paragraph_format.keep_with_next = True

    # footer page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")

    # cover
    for _ in range(7):
        doc.add_paragraph()
    t = doc.add_paragraph()
    r = t.add_run(TITLE)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(28), RGBColor(0x12, 0x35, 0x5B)
    for line in (SUBTITLE, DATE):
        p = doc.add_paragraph()
        rr = p.add_run(line)
        rr.font.size, rr.font.color.rgb = Pt(13), RGBColor(0x44, 0x44, 0x44)
    doc.add_page_break()

    # TOC (static, internally hyperlinked)
    hs = [(l, tx) for l, tx in headings(combined) if l <= 2 or (l == 3 and re.match(r"Week \d+", tx))]
    toc_h = doc.add_paragraph(style="Heading 1")
    toc_h.add_run("Contents")
    note = doc.add_paragraph()
    nr = note.add_run("Click an entry to jump. In Word, the Navigation Pane (View > Navigation Pane) shows every heading.")
    nr.italic, nr.font.size = True, Pt(9)
    anchors = {}
    for i, (lvl, tx) in enumerate(hs):
        anchors[(lvl, tx, i)] = f"_toc{i}"
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
        p.paragraph_format.space_after = Pt(1 if lvl > 1 else 2)
        p.paragraph_format.space_before = Pt(5 if lvl == 1 else 0)
        internal_link(p, f"_toc{i}", tx, 10.5 if lvl == 1 else 9)
        if lvl == 1:
            for run in p._p.iter(qn("w:r")):
                rp = run.get_or_add_rPr() if hasattr(run, "get_or_add_rPr") else None
    heading_cursor = 0

    # body
    lines = combined.split("\n")
    i, n, bm_id = 0, len(lines), 100
    in_code, code_buf = False, []
    first_h1 = True
    while i < n:
        ln = lines[i]
        if FENCE.match(ln):
            if in_code:
                p = doc.add_paragraph()
                pPr = p._p.get_or_add_pPr()
                shade(pPr, "F3F5F7")
                p.paragraph_format.left_indent = Cm(0.3)
                for j, cl in enumerate(code_buf):
                    run = p.add_run(cl if cl else " ")
                    run.font.name, run.font.size = "Menlo", Pt(8.5)
                    if j < len(code_buf) - 1:
                        run.add_break()
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(ln.rstrip("\n"))
            i += 1
            continue
        if not ln.strip():
            i += 1
            continue
        hm = HEAD.match(ln)
        if hm:
            lvl, tx = len(hm.group(1)), hm.group(2)
            lvl = min(lvl, 4)
            if lvl == 1:
                doc.add_page_break()
            p = doc.add_paragraph(style=f"Heading {lvl}")
            add_inline(p, tx)
            # `hs` was built with this same filter and order, so the Nth TOC heading gets bookmark _toc{N}
            raw = len(hm.group(1))
            if raw <= 2 or (raw == 3 and re.match(r"Week \d+", tx)):
                bm_id += 1
                bookmark(p, f"_toc{heading_cursor}", bm_id)
                heading_cursor += 1
            i += 1
            continue
        if ln.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = [split_row(ln)]
            i += 2
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            ncol = max(len(r) for r in rows)
            tb = doc.add_table(rows=len(rows), cols=ncol)
            tb.style = "Table Grid"
            tb.autofit = True
            fs = 8 if ncol >= 5 else 9
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = tb.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(1)
                    txt = row[ci] if ci < len(row) else ""
                    if ri == 0:
                        set_cell_bg(cell, "12355B")
                        add_inline(par, txt, fs, base_bold=True, color=RGBColor(255, 255, 255))
                    else:
                        if ri % 2 == 0:
                            set_cell_bg(cell, "F4F7FA")
                        add_inline(par, txt, fs)
            # repeat header row
            trPr = tb.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trPr.append(th)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if ln.startswith(">"):
            buf = []
            while i < n and lines[i].startswith(">"):
                buf.append(lines[i].lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            shade(p._p.get_or_add_pPr(), "FFF7E0")
            par_border_left(p)
            p.paragraph_format.left_indent = Cm(0.4)
            add_inline(p, " ".join(b for b in buf if b))
            continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if lm:
            indent = len(lm.group(1))
            marker, body = lm.group(2), lm.group(3)
            deep = indent >= 2
            if marker in "-*":
                style = "List Bullet 2" if deep else "List Bullet"
                p = doc.add_paragraph(style=style)
                if body.startswith(("☐", "☑")):
                    # checkbox: keep the glyph, remove bullet look by using plain indent
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(1.0 if deep else 0.5)
                    p.paragraph_format.first_line_indent = Cm(-0.5)
                add_inline(p, body)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.4 if deep else 0.9)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                add_inline(p, f"{marker} {body}")
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        # paragraph (merge soft-wrapped lines)
        buf = [ln.strip()]
        i += 1
        while i < n and lines[i].strip() and not (HEAD.match(lines[i]) or FENCE.match(lines[i]) or lines[i].startswith(">")
                                                    or lines[i].lstrip().startswith("|") or re.match(r"^\s*([-*]|\d+\.)\s+", lines[i])):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))
    doc.save(OUT_DOCX)


def main():
    combined = normalise_checks(combine())
    OUT_MD.write_text(combined, encoding="utf-8")
    build_docx(combined)
    print("DOCX:", OUT_DOCX, OUT_DOCX.stat().st_size // 1024, "KB")
    if build_pdf(combined):
        print("PDF:", OUT_PDF, OUT_PDF.stat().st_size // 1024, "KB")
    print("words:", len(combined.split()))


if __name__ == "__main__":
    sys.exit(main())
